"""Shared CSV/Excel import helpers for the tracker routers.

Companies and contacts both accept CSV or Excel (.xlsx/.xlsm) uploads with
flexible header aliases and a per-row error report. Everything generic lives
here; each router supplies its own field->aliases map and import field list.
"""

import csv
import io
import json
import logging
from typing import Any

from fastapi import HTTPException, UploadFile

logger = logging.getLogger(__name__)

# Maximum accepted import file size (bytes).
MAX_IMPORT_BYTES = 5 * 1024 * 1024


def normalize_header(value: str) -> str:
    """Normalize a header/field value for alias matching."""
    return (
        value.strip().lower().replace(" ", "").replace("_", "").replace("-", "").replace("'", "")
    )


def detect_delimiter(sample: str) -> str:
    """Pick the CSV delimiter, handling Excel locale exports (`;`, tab, `|`).

    Falls back to comma when the sample is a single column or unparseable.
    """
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        return ","


def rows_from_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig", errors="replace")
    delimiter = detect_delimiter(text[:4096])
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [dict(row) for row in reader]


def rows_from_xlsx(content: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers: list[str] = []
    records: list[dict[str, str]] = []
    for row in rows:
        if not headers:
            headers = [str(cell) if cell is not None else "" for cell in row]
            continue
        records.append(
            {header: (str(cell) if cell is not None else "") for header, cell in zip(headers, row)}
        )
    return records


def normalize_record(
    record: dict[str, str],
    aliases: dict[str, set[str]],
    mapping: dict[str, str] | None = None,
) -> dict[str, str]:
    """Map a raw CSV/Excel row's headers to our field names (empty = unset).

    With a ``mapping`` ({field: exact header}) only the mapped columns are
    used; without one, headers are matched through ``aliases``.
    """
    lookup: dict[str, str] = {}
    for header, value in record.items():
        if header is None:
            continue
        lookup[normalize_header(str(header))] = value
    normalized: dict[str, str] = {}
    if mapping is not None:
        for field, header in mapping.items():
            if field not in aliases or not header:
                continue
            value = lookup.get(normalize_header(header))
            if value is not None:
                normalized[field] = value
        return normalized
    for field, field_aliases in aliases.items():
        for alias in field_aliases:
            if alias in lookup:
                normalized[field] = lookup[alias]
                break
    return normalized


def read_rows(content: bytes, filename: str) -> list[dict[str, str]]:
    """Parse CSV/Excel bytes into raw header->value records."""
    if filename.endswith(".csv"):
        return rows_from_csv(content)
    if filename.endswith((".xlsx", ".xlsm")):
        try:
            return rows_from_xlsx(content)
        except Exception as e:
            logger.warning("Failed to parse Excel import: %s", e)
            raise HTTPException(status_code=422, detail="Could not parse the Excel file.")
    raise HTTPException(
        status_code=422, detail="Unsupported file type — upload a .csv or .xlsx file."
    )


async def read_upload(file: UploadFile) -> tuple[list[dict[str, str]], str]:
    """Validate and parse an import upload, returning (rows, lowercased filename)."""
    filename = (file.filename or "").lower()
    content = await file.read()
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Import file is too large (max 5 MB).")
    if not content:
        raise HTTPException(status_code=422, detail="Import file is empty.")
    return read_rows(content, filename), filename


def parse_mapping(raw: str | None, aliases: dict[str, set[str]]) -> dict[str, str] | None:
    """Parse the optional ``mapping`` form field into {field: header}."""
    if raw is None or not raw.strip():
        return None
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        raise HTTPException(status_code=422, detail="Mapping must be a valid JSON object.")
    if not isinstance(data, dict):
        raise HTTPException(status_code=422, detail="Mapping must be a JSON object.")
    mapping: dict[str, str] = {}
    for field, header in data.items():
        if field not in aliases:
            continue
        mapping[field] = header if isinstance(header, str) else ""
    return mapping


def detect_mapped_headers(
    headers: list[str], aliases: dict[str, set[str]], import_fields: list[str]
) -> dict[str, str]:
    """Pick which import fields each file header already matches."""
    detected: dict[str, str] = {}
    lookup = {normalize_header(h): h for h in headers}
    for field in import_fields:
        for alias in aliases[field]:
            if alias in lookup:
                detected[field] = lookup[alias]
                break
    return detected
